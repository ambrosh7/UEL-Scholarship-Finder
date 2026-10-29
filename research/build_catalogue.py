#!/usr/bin/env python3
"""Turn Scrapy + live-check notes into catalogue JSON, sources.csv and Task 2 xlsx."""

from __future__ import annotations

import csv
import json
from pathlib import Path

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parents[1]
CATALOGUE = ROOT / "data" / "catalogue"
CHECKED_ON = "2026-10-28"

INDIA = "https://www.uel.ac.uk/international/guidance-your-region/south-asia/india"
SCHOLARSHIP = "https://www.uel.ac.uk/study/fees-funding/international-fees-funding/international-scholarships"
FEES_PDF = "https://www.uel.ac.uk/sites/default/files/uel-fees-schedule-2026-27.pdf"
INTL_FEES = "https://www.uel.ac.uk/study/fees-funding/international-fees-funding"
DEPOSIT = "https://www.uel.ac.uk/study/fees-funding/paying-your-fees-international-students"
ACCOM = "https://www.uel.ac.uk/study/accommodation/room-types-prices-payment-plans"
RBI = "https://www.rbi.org.in/"
GOV_VISA = "https://www.gov.uk/student-visa"
GOV_IHS = "https://www.gov.uk/healthcare-immigration-application"
GOV_MONEY = "https://www.gov.uk/student-visa/money"

UG_INDIA = {
    "qualification": "class-12",
    "boards": ["CBSE", "ISC", "CICSE", "AISSCE"],
    "minPercent": 60,
    "notes": (
        "India page, live-checked 2026-10-28: HSC / Standard 12 overall 60% "
        "(CBSE/CICSE and named state boards). Kerala State Board 65% in Year 12. "
        "All other boards 70%. Diploma 55% is a separate route, not used as the Class 12 threshold."
    ),
}

PG_INDIA = {
    "qualification": "bachelor",
    "minPercent": 50,
    "notes": (
        "India page, live-checked 2026-10-28: 50% from listed universities "
        "(Delhi, Mumbai, Pune, Kolkata, Chennai, Anna, Bharatiyar, MSU Baroda, Bangalore); "
        "55% from all other recognised institutions."
    ),
}

YEAR_START_ISSUE = (
    "JSON-LD start dates on the course page are 2027 (and some 2026-09 / 2027-05 for PGT). "
    "Checked 2026-10-28, after the September 2026 teaching start. Fees still match the 2026/27 schedule."
)

TBC_UG_FEE_ISSUE = (
    "Course page international fee is TBC. Figure taken from UEL Fees Schedule 2026/27 "
    "international undergraduate headline £16,020 (all schools except listed nursing exceptions)."
)

COURSES = [
    {
        "id": "bsc-business-management",
        "title": "Business Management BSc (Hons)",
        "level": "undergraduate",
        "subjectFamily": "business",
        "campus": "University Square Stratford",
        "durationYears": 3,
        "tuitionGbpYearOne": 16020,
        "tuitionSource": "course-page",
        "startDates": ["2027-01", "2027-09"],
        "extraConditions": [
            "IELTS 5.5 with a minimum of 5.5 in Reading, Writing, Listening and Speaking (or recognised equivalent).",
        ],
        "sourceUrl": "https://www.uel.ac.uk/undergraduate/courses/bsc-hons-business-management",
        "liveChecked": True,
        "issues": [YEAR_START_ISSUE],
    },
    {
        "id": "bsc-marketing",
        "title": "Marketing BSc (Hons)",
        "level": "undergraduate",
        "subjectFamily": "business",
        "campus": "University Square Stratford",
        "durationYears": 3,
        "tuitionGbpYearOne": 16020,
        "tuitionSource": "course-page",
        "startDates": ["2027-01", "2027-09"],
        "extraConditions": [
            "IELTS 5.5 with a minimum of 5.5 in Reading, Writing, Listening and Speaking (or recognised equivalent).",
            "Page notes CIM professional accreditation; direct entry students only considered for Level 5 admission.",
        ],
        "sourceUrl": "https://www.uel.ac.uk/undergraduate/courses/bsc-hons-marketing",
        "issues": [YEAR_START_ISSUE],
    },
    {
        "id": "bsc-accounting-finance",
        "title": "Accounting and Finance BSc (Hons)",
        "level": "undergraduate",
        "subjectFamily": "business",
        "campus": "University Square Stratford",
        "durationYears": 3,
        "tuitionGbpYearOne": 16020,
        "tuitionSource": "course-page",
        "startDates": ["2027-01", "2027-09"],
        "extraConditions": [
            "IELTS 5.5 with a minimum of 5.5 in Reading, Writing, Listening and Speaking (or recognised equivalent).",
        ],
        "sourceUrl": "https://www.uel.ac.uk/undergraduate/courses/bsc-hons-accounting-finance",
        "issues": [YEAR_START_ISSUE],
    },
    {
        "id": "bsc-accounting-business-management",
        "title": "Accounting and Business Management BSc (Hons)",
        "level": "undergraduate",
        "subjectFamily": "business",
        "campus": "University Square Stratford",
        "durationYears": 3,
        "tuitionGbpYearOne": 16020,
        "tuitionSource": "fees-schedule",
        "startDates": [],
        "extraConditions": [
            "IELTS 5.5 with a minimum of 5.5 in Reading, Writing, Listening and Speaking (or recognised equivalent).",
            "Page states the course is subject to revalidation.",
        ],
        "sourceUrl": "https://www.uel.ac.uk/undergraduate/courses/bsc-hons-accounting-business-management",
        "issues": [
            "Swapped in for International Business UG: UEL publishes International Business Management as a one-year top-up, not a full bachelor's.",
            TBC_UG_FEE_ISSUE,
            "JSON-LD on this page did not list international start dates. Not invented.",
        ],
    },
    {
        "id": "bsc-business-economics",
        "title": "Business Economics BSc (Hons)",
        "level": "undergraduate",
        "subjectFamily": "business",
        "campus": "University Square Stratford",
        "durationYears": 3,
        "tuitionGbpYearOne": 16020,
        "tuitionSource": "fees-schedule",
        "startDates": [],
        "extraConditions": [
            "IELTS 5.5 with a minimum of 5.5 in Reading, Writing, Listening and Speaking (or recognised equivalent).",
        ],
        "sourceUrl": "https://www.uel.ac.uk/undergraduate/courses/bsc-hons-business-economics",
        "issues": [
            TBC_UG_FEE_ISSUE,
            "JSON-LD on this page did not list international start dates. Not invented.",
        ],
    },
    {
        "id": "msc-international-business-management",
        "title": "International Business Management MSc",
        "level": "postgraduate-taught",
        "subjectFamily": "business",
        "campus": "University Square Stratford",
        "durationYears": 1,
        "tuitionGbpYearOne": 17700,
        "tuitionSource": "course-page",
        "startDates": ["2026-09", "2027-01", "2027-05"],
        "extraConditions": [
            "IELTS 6.0 with a minimum of 6.0 in Writing, Speaking, Reading and Listening (or recognised equivalent).",
            "Page states you do not need a business background to apply.",
        ],
        "sourceUrl": "https://www.uel.ac.uk/postgraduate/courses/msc-international-business-management",
        "liveChecked": True,
        "issues": [
            YEAR_START_ISSUE,
            "Campus section describes both Docklands and University Square Stratford. USS is stored as the RDSBL/MBA campus named on the page; Docklands is also mentioned.",
        ],
    },
    {
        "id": "msc-marketing",
        "title": "Marketing MSc",
        "level": "postgraduate-taught",
        "subjectFamily": "business",
        "campus": "University Square Stratford",
        "durationYears": 1,
        "tuitionGbpYearOne": 18900,
        "tuitionSource": "course-page",
        "startDates": ["2026-09", "2027-01"],
        "extraConditions": [
            "IELTS score of 6 (or recognised equivalent).",
        ],
        "sourceUrl": "https://www.uel.ac.uk/postgraduate/courses/msc-marketing",
        "issues": [YEAR_START_ISSUE],
    },
    {
        "id": "mba",
        "title": "Master of Business Administration MBA",
        "level": "postgraduate-taught",
        "subjectFamily": "business",
        "campus": "University Square Stratford",
        "durationYears": 1,
        "tuitionGbpYearOne": 19500,
        "tuitionSource": "course-page",
        "startDates": ["2026-09", "2027-01", "2027-05"],
        "extraConditions": [
            "Bachelor's degree with minimum Second Class (2:2) or equivalent in any subject.",
            "IELTS 6.0 with a minimum of 6.0 in Writing and Speaking; minimum 5.5 in Reading and Listening (or recognised equivalent).",
        ],
        "sourceUrl": "https://www.uel.ac.uk/postgraduate/courses/mba-master-business-administration",
        "issues": [
            YEAR_START_ISSUE,
            "Page says USS is the base for the MBA and also describes an MBA suite at Docklands. Stored campus is University Square Stratford.",
        ],
    },
    {
        "id": "msc-human-resource-management",
        "title": "Human Resource Management MSc",
        "level": "postgraduate-taught",
        "subjectFamily": "business",
        "campus": "University Square Stratford",
        "durationYears": 1,
        "tuitionGbpYearOne": 17220,
        "tuitionSource": "course-page",
        "startDates": ["2026-09", "2027-01", "2027-05"],
        "extraConditions": [
            "Bachelor's degree with minimum Second Class (2:2) or equivalent in any subject.",
            "IELTS 6.0 with a minimum of 6.0 in Writing and Speaking; minimum 5.5 in Reading and Listening (or recognised equivalent).",
        ],
        "sourceUrl": "https://www.uel.ac.uk/postgraduate/courses/msc-human-resource-management",
        "issues": [YEAR_START_ISSUE],
    },
    {
        "id": "bsc-computer-science",
        "title": "Computer Science BSc (Hons)",
        "level": "undergraduate",
        "subjectFamily": "computing",
        "campus": "Docklands",
        "durationYears": 3,
        "tuitionGbpYearOne": 16020,
        "tuitionSource": "course-page",
        "startDates": ["2027-01", "2027-09"],
        "extraConditions": [
            "IELTS 5.5 with a minimum of 5.5 in Reading, Writing, Listening and Speaking (or recognised equivalent).",
            "UK IB tariff on the page requires Mathematics at Higher Level. India page: Year 12 Maths 50% only if the course requires maths.",
        ],
        "sourceUrl": "https://www.uel.ac.uk/undergraduate/courses/bsc-hons-computer-science",
        "liveChecked": True,
        "issues": [YEAR_START_ISSUE],
    },
    {
        "id": "msc-computer-science",
        "title": "Computer Science MSc",
        "level": "postgraduate-taught",
        "subjectFamily": "computing",
        "campus": "Docklands",
        "durationYears": 1,
        "tuitionGbpYearOne": 17220,
        "tuitionSource": "course-page",
        "startDates": ["2026-09", "2027-01", "2027-05"],
        "extraConditions": [
            "Bachelor's degree with minimum Second Class (2:2) in Computer Science (or equivalent degree with significant maths and computing content).",
            "IELTS 6.0 with minimum 6.0 in writing and speaking, and 5.5 in listening and reading (or recognised equivalent).",
        ],
        "sourceUrl": "https://www.uel.ac.uk/postgraduate/courses/msc-computer-science",
        "issues": [YEAR_START_ISSUE],
    },
    {
        "id": "msc-data-science",
        "title": "Data Science MSc",
        "level": "postgraduate-taught",
        "subjectFamily": "computing",
        "campus": "Docklands",
        "durationYears": 1,
        "tuitionGbpYearOne": 17220,
        "tuitionSource": "course-page",
        "startDates": ["2026-09", "2027-01"],
        "extraConditions": [
            "Bachelor's degree with minimum Second Class (2:2) or equivalent in any subject.",
            "IELTS 6.0 with minimum 6.0 in writing and speaking, and 5.5 in listening and reading (or recognised equivalent).",
            "Page states this is not a computer science programme; it is methodology-oriented.",
        ],
        "sourceUrl": "https://www.uel.ac.uk/postgraduate/courses/msc-data-science",
        "issues": [
            YEAR_START_ISSUE,
            "Any-subject 2:2 on this MSc is different from Computer Science MSc, which asks for a computing degree.",
        ],
    },
    {
        "id": "beng-civil-engineering",
        "title": "Civil Engineering BEng (Hons)",
        "level": "undergraduate",
        "subjectFamily": "engineering",
        "campus": "Docklands",
        "durationYears": 3,
        "tuitionGbpYearOne": 16020,
        "tuitionSource": "course-page",
        "startDates": ["2027-01", "2027-09"],
        "extraConditions": [
            "A Levels must include Mathematics grade C. India page if maths required: HSC Mathematics (Year 12) 50% or higher; Year 10 80%+ if Maths was not taken in Year 12.",
            "IELTS 6.0 with a minimum of 6.0 in Writing and Speaking; 5.5 in Reading and Listening (or recognised equivalent).",
        ],
        "sourceUrl": "https://www.uel.ac.uk/undergraduate/courses/beng-hons-civil-engineering",
        "issues": [YEAR_START_ISSUE],
        "indiaSubjectConditions": [
            "HSC Mathematics (Year 12) with 50% or higher (course requires maths).",
        ],
    },
    {
        "id": "beng-mechanical-engineering",
        "title": "Mechanical Engineering BEng (Hons)",
        "level": "undergraduate",
        "subjectFamily": "engineering",
        "campus": "Docklands",
        "durationYears": 3,
        "tuitionGbpYearOne": 16020,
        "tuitionSource": "fees-schedule",
        "startDates": ["2027-09"],
        "extraConditions": [
            "A Levels must include Mathematics and Physics or Electronics. India page if maths required: HSC Mathematics (Year 12) 50% or higher.",
            "IELTS 6.0 with a minimum of 6.0 in Writing and Speaking; 5.5 in Listening and Reading (or recognised equivalent).",
        ],
        "sourceUrl": "https://www.uel.ac.uk/undergraduate/courses/beng-hons-mechanical-engineering",
        "issues": [YEAR_START_ISSUE, TBC_UG_FEE_ISSUE],
        "indiaSubjectConditions": [
            "HSC Mathematics (Year 12) with 50% or higher (course requires maths).",
        ],
    },
    {
        "id": "bsc-psychology",
        "title": "Psychology BSc (Hons)",
        "level": "undergraduate",
        "subjectFamily": "other",
        "campus": "Stratford",
        "durationYears": 3,
        "tuitionGbpYearOne": 16020,
        "tuitionSource": "fees-schedule",
        "startDates": ["2027-09"],
        "extraConditions": [
            "IELTS 6.0 with a minimum of 6.0 in Writing and Speaking; 5.5 in Listening and Reading (or recognised equivalent).",
            "Page describes the programme as BPS-accredited.",
        ],
        "sourceUrl": "https://www.uel.ac.uk/undergraduate/courses/bsc-hons-psychology",
        "issues": [YEAR_START_ISSUE, TBC_UG_FEE_ISSUE],
    },
    {
        "id": "llb-law",
        "title": "Law LLB (Hons)",
        "level": "undergraduate",
        "subjectFamily": "law",
        "campus": "University Square Stratford",
        "durationYears": 3,
        "tuitionGbpYearOne": 16020,
        "tuitionSource": "course-page",
        "startDates": ["2027-01", "2027-09"],
        "extraConditions": [
            "IELTS 6.0 with a minimum of 6.0 in Writing and Speaking; 5.5 in Listening and Reading (or recognised equivalent).",
        ],
        "sourceUrl": "https://www.uel.ac.uk/undergraduate/courses/llb-hons-law",
        "issues": [YEAR_START_ISSUE],
    },
]


def catalogue_courses() -> list[dict]:
    rows = []
    for course in COURSES:
        india = dict(UG_INDIA if course["level"] == "undergraduate" else PG_INDIA)
        if course.get("indiaSubjectConditions"):
            india["subjectConditions"] = course["indiaSubjectConditions"]
        rows.append(
            {
                "id": course["id"],
                "title": course["title"],
                "level": course["level"],
                "subjectFamily": course["subjectFamily"],
                "campus": course["campus"],
                "durationYears": course["durationYears"],
                "tuitionGbpYearOne": course["tuitionGbpYearOne"],
                "startDates": course["startDates"],
                "indiaEntry": india,
                "extraConditions": course["extraConditions"],
                "scholarshipEligible": True,
                "sourceUrl": course["sourceUrl"],
                "checkedOn": CHECKED_ON,
                "issues": course["issues"],
            }
        )
    return rows


def scholarships() -> list[dict]:
    return [
        {
            "id": "uel-international-scholarship",
            "amountMinGbp": 500,
            "amountMaxGbp": 5000,
            "appliesTo": "first-year-tuition",
            "autoAssessed": True,
            "requiresOffer": True,
            "eligibleLevels": ["undergraduate", "postgraduate-taught"],
            "exclusions": [
                "Short courses or exchange/study abroad programmes",
                "Franchised programmes",
                "Pre-Masters programmes",
                "International Foundation programmes",
                "Pre-sessional English programmes",
            ],
            "sourceUrl": SCHOLARSHIP,
            "checkedOn": CHECKED_ON,
        }
    ]


def costs() -> dict:
    return {
        "academicYear": "2026/27",
        "londonLivingGbp": "PLACEHOLDER",
        "visaGbp": "PLACEHOLDER",
        "ihsGbp": "PLACEHOLDER",
        "depositGbp": 1500,
        "accommodationGbp": 8010,
        "flightsGbp": "PLACEHOLDER",
        "inrPerGbp": 126.4657,
        "inrSourceUrl": RBI,
        "notes": [
            "Visa, IHS and London living/maintenance are PLACEHOLDER until Ambrosh spot-checks GOV.UK.",
            "Tuition deposit: £1,500 starts pre-CAS checks; the full deposit is in the offer letter and a further payment is required for CAS. Source: paying-your-fees-international-students, 2026-10-28.",
            "Typical room: East standard full academic year 12/09/2026–06/06/2027, £8,010 (£210/week). Range on the same page: West standard £182/week in the room table vs intro/FAQ 'from £189/week'. Halls security deposit £250.",
            "Flights India → London not sourced. PLACEHOLDER.",
            "INR snapshot: RBI reference INR 126.4657 per GBP as at 1.00pm 28 Oct 2026 (source FBIL on the RBI homepage). Not live FX.",
        ],
        "sourceUrl": INTL_FEES,
        "checkedOn": CHECKED_ON,
    }


def source_rows(courses: list[dict]) -> list[dict]:
    rows: list[dict] = []

    def add(field, value, url, collector, issues=""):
        rows.append(
            {
                "field": field,
                "value": value,
                "url": url,
                "checkedOn": CHECKED_ON,
                "collector": collector,
                "issues": issues,
            }
        )

    add(
        "robots.txt",
        "Course, scholarship and India URLs are allowed; /admin /user/login /search /node disallowed. No login used.",
        "https://www.uel.ac.uk/robots.txt",
        "scrapy",
        "",
    )

    for course in COURSES:
        collector = "live-check" if course.get("liveChecked") else "scrapy"
        fee_url = FEES_PDF if course["tuitionSource"] == "fees-schedule" else course["sourceUrl"]
        fee_issue = TBC_UG_FEE_ISSUE if course["tuitionSource"] == "fees-schedule" else ""
        add(f"{course['id']}.title", course["title"], course["sourceUrl"], collector, "")
        add(
            f"{course['id']}.tuitionGbpYearOne",
            str(course["tuitionGbpYearOne"]),
            fee_url,
            "live-check" if course["tuitionSource"] == "fees-schedule" or course.get("liveChecked") else "scrapy",
            fee_issue,
        )
        add(f"{course['id']}.campus", course["campus"], course["sourceUrl"], collector, "; ".join(course["issues"]))
        add(
            f"{course['id']}.startDates",
            ",".join(course["startDates"]) if course["startDates"] else "PLACEHOLDER",
            course["sourceUrl"],
            collector,
            "No international start dates in JSON-LD." if not course["startDates"] else YEAR_START_ISSUE,
        )
        add(f"{course['id']}.durationYears", str(course["durationYears"]), course["sourceUrl"], collector, "")
        india = UG_INDIA if course["level"] == "undergraduate" else PG_INDIA
        add(
            f"{course['id']}.indiaEntry.minPercent",
            str(india["minPercent"]),
            INDIA,
            "live-check",
            india["notes"],
        )

    add("scholarship.amountMinGbp", "500", SCHOLARSHIP, "live-check", "")
    add("scholarship.amountMaxGbp", "5000", SCHOLARSHIP, "live-check", "")
    add("scholarship.appliesTo", "first-year-tuition", SCHOLARSHIP, "live-check", "")
    add("scholarship.autoAssessed", "true", SCHOLARSHIP, "live-check", "No separate application. Shown in the offer letter.")
    add(
        "scholarship.exclusions",
        "Short courses / exchange; franchised; Pre-Masters; International Foundation; pre-sessional English",
        SCHOLARSHIP,
        "live-check",
        "",
    )
    add("costs.depositGbp", "1500", DEPOSIT, "live-check", "£1,500 starts pre-CAS; full deposit is in the offer letter.")
    add(
        "costs.accommodationGbp",
        "8010",
        ACCOM,
        "live-check",
        "East standard full academic year 2026/27. Intro/FAQ says from £189/week; West standard table shows £182/week.",
    )
    add("costs.accommodationDepositGbp", "250", "https://www.uel.ac.uk/study/accommodation/how-apply-student-accommodation-uel-202627", "scrapy", "Refundable halls security deposit.")
    add("costs.londonLivingGbp", "PLACEHOLDER", GOV_MONEY, "govuk-spot-check", "Awaiting Ambrosh GOV.UK spot-check.")
    add("costs.visaGbp", "PLACEHOLDER", GOV_VISA, "govuk-spot-check", "Awaiting Ambrosh GOV.UK spot-check.")
    add("costs.ihsGbp", "PLACEHOLDER", GOV_IHS, "govuk-spot-check", "Awaiting Ambrosh GOV.UK spot-check.")
    add("costs.flightsGbp", "PLACEHOLDER", "PLACEHOLDER", "scrapy", "Not sourced. Do not invent a fare.")
    add("costs.inrPerGbp", "126.4657", RBI, "rbi", "As at 1.00pm 28 Oct 2026. Source FBIL, displayed on RBI homepage.")
    add("ug.internationalHeadlineGbp", "16020", FEES_PDF, "live-check", "2026/27 international UG headline, all schools except listed nursing exceptions.")
    add("mba.internationalGbp", "19500", INTL_FEES, "live-check", "International fees page and MBA course page agree.")
    return rows


def write_json() -> tuple[list[dict], list[dict], dict]:
    CATALOGUE.mkdir(parents=True, exist_ok=True)
    courses = catalogue_courses()
    policies = scholarships()
    living = costs()
    (CATALOGUE / "courses.json").write_text(json.dumps(courses, indent=2, ensure_ascii=False) + "\n")
    (CATALOGUE / "scholarships.json").write_text(json.dumps(policies, indent=2, ensure_ascii=False) + "\n")
    (CATALOGUE / "costs.json").write_text(json.dumps(living, indent=2, ensure_ascii=False) + "\n")
    return courses, policies, living


def write_csv(rows: list[dict]) -> None:
    path = CATALOGUE / "sources.csv"
    with path.open("w", newline="", encoding="utf-8") as handle:
        writer = csv.DictWriter(
            handle,
            fieldnames=["field", "value", "url", "checkedOn", "collector", "issues"],
        )
        writer.writeheader()
        writer.writerows(rows)


def write_note() -> None:
    text = f"""# Task 2 collection note

Checked on {CHECKED_ON}. Academic year **2026/27**.

## How we collected it

1. Confirmed [UEL robots.txt](https://www.uel.ac.uk/robots.txt). Course, scholarship and India URLs are allowed. `/admin`, `/user/login`, `/search` and `/node` are disallowed. We never logged in.
2. **Scrapy first.** Polite crawl (`DOWNLOAD_DELAY=3`, one request at a time, custom user-agent, cookies off) of 16 course pages plus scholarships, India entry, international fees, deposit, accommodation, admission deadlines and robots.txt. Raw extract and HTML: `data/raw/`.
3. **Then every UEL fact was checked against the live page HTML** from that crawl. Browser live-check of Business Management BSc, Computer Science BSc, International Business Management MSc, the international scholarships page, and the India page.
4. **GOV.UK** student visa, Immigration Health Surcharge and London living / maintenance are **PLACEHOLDER** until Ambrosh spot-checks them. Rows are in `sources.csv` with collector `govuk-spot-check`.
5. **RBI** INR/GBP reference rate taken from [rbi.org.in](https://www.rbi.org.in/) as at 1.00pm 28 October 2026 (quoted source FBIL): **126.4657**. Labelled snapshot, not live FX.

## Basket changes

- There is no full-length International Business bachelor's. UEL publishes a one-year **International Business Management Top-up**. Swapped to **Accounting and Business Management BSc (Hons)** and added **Business Economics BSc (Hons)** to keep the mix management-heavy.
- Extra management master's: **Human Resource Management MSc**.
- Computing master's: both **Computer Science MSc** and **Data Science MSc** (the Data Science page says it is not a computer science programme).

## Issues found

- Several course pages list international fees as **TBC** while the 2026/27 fees schedule still publishes a £16,020 undergraduate international headline. Those courses store the schedule figure and say so in `issues`.
- Course JSON-LD start dates are mostly **2027**. Checked after September 2026 teaching had started. Stored as published; not rewritten to 2026-09.
- MBA / International Business MSc pages describe **both Docklands and University Square Stratford**.
- Accommodation intro/FAQ says from **£189/week**; the West standard table shows **£182/week**. Typical room stored is East standard full year **£8,010**.
- India page staff bio contained stray copy (“this is mine”) on the live page. Not used as a fact.
- Flights India → London: unsourced, **PLACEHOLDER**.

Scrapy does not replace the live check. `data/raw` is the scrape; `data/catalogue` is the checked JSON the app will read.
"""
    (CATALOGUE / "collection-note.md").write_text(text)


def write_xlsx(courses: list[dict], rows: list[dict], living: dict) -> None:
    wb = Workbook()
    arial = Font(name="Arial", size=11)
    arial_bold = Font(name="Arial", size=11, bold=True)
    arial_title = Font(name="Arial", size=16, bold=True)
    blue = Font(name="Arial", size=11, color="0000FF")
    green = Font(name="Arial", size=11, color="008000")
    yellow = PatternFill("solid", fgColor="FFFF00")
    header_fill = PatternFill("solid", fgColor="D9E2F3")
    thin = Border(
        left=Side(style="thin", color="B0B0B0"),
        right=Side(style="thin", color="B0B0B0"),
        top=Side(style="thin", color="B0B0B0"),
        bottom=Side(style="thin", color="B0B0B0"),
    )
    wrap = Alignment(wrap_text=True, vertical="top")

    def style_header(ws, width_map):
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
        for cell in ws[1]:
            cell.font = Font(name="Arial", size=11, bold=True, color="1F4E79")
            cell.fill = header_fill
            cell.alignment = wrap
            cell.border = thin
        for col, width in width_map.items():
            ws.column_dimensions[col].width = width
        ws.row_dimensions[1].height = 22
        for row in ws.iter_rows(min_row=2, max_row=ws.max_row):
            ws.row_dimensions[row[0].row].height = 36
            for cell in row:
                cell.font = arial
                cell.alignment = wrap
                cell.border = thin

    # Cover
    cover = wb.active
    cover.title = "Cover"
    cover["A1"] = "UEL Fit & Scholarship Finder — Task 2 research"
    cover["A1"].font = arial_title
    cover.merge_cells("A1:F1")
    cover["A3"] = "Academic year"
    cover["B3"] = "2026/27"
    cover["A4"] = "Checked on"
    cover["B4"] = CHECKED_ON
    cover["A5"] = "Method"
    cover["B5"] = "Scrapy → live UEL check → Ambrosh GOV.UK check → RBI snapshot"
    cover["A7"] = "Counts (formulas — update when the other tabs change)"
    cover["A7"].font = arial_bold
    cover["A8"] = "Courses in catalogue"
    cover["B8"] = "=COUNTA(Courses!A3:A30)"
    cover["A9"] = "Source rows"
    cover["B9"] = '=COUNTA(Sources!A:A)-1'
    cover["A10"] = "PLACEHOLDER cost lines"
    cover["B10"] = '=COUNTIF(Costs!C2:C7,"PLACEHOLDER")'
    cover["A11"] = "Management courses"
    cover["B11"] = '=COUNTIF(Courses!D3:D30,"business")'
    cover["C8"] = "Must be at least 15"
    cover["C10"] = "Visa, IHS, London living, flights until Ambrosh / a sourced fare"

    cover["A13"] = "What to fill"
    cover["A13"].font = arial_bold
    cover["A14"] = (
        "Yellow cells on the GOV.UK tab are for Ambrosh: visa fee, IHS, and London living / maintenance, "
        "each with URL and date. Do not type invented numbers into the catalogue."
    )
    cover.merge_cells("A14:F16")
    cover["A14"].alignment = wrap
    cover["A18"] = "Collection sequence"
    cover["A18"].font = arial_bold
    cover["A19"] = (
        "1. robots.txt — course/scholarship/India paths allowed; no login.\n"
        "2. Scrapy polite crawl (3s delay, 1 concurrent) into data/raw.\n"
        "3. Live check of UEL facts against the page (browser sample: Business Management BSc, "
        "Computer Science BSc, International Business Management MSc, scholarships, India).\n"
        "4. GOV.UK visa / IHS / living — Ambrosh.\n"
        "5. RBI INR/GBP reference from rbi.org.in, 1.00pm 28 Oct 2026, 126.4657 (FBIL)."
    )
    cover.merge_cells("A19:F24")
    cover["A19"].alignment = wrap
    cover["A26"] = "Blue text = sourced input copied from a page. Green text = formula. Yellow = Ambrosh to complete."
    cover["A26"].font = Font(name="Arial", size=10, italic=True)
    for cell in ["A3", "A4", "A5", "A8", "A9", "A10", "A11"]:
        cover[cell].font = arial_bold
    for cell in ["B3", "B4", "B5"]:
        cover[cell].font = blue
    for cell in ["B8", "B9", "B10", "B11"]:
        cover[cell].font = Font(name="Arial", size=11, color="008000")
        cover[cell].comment = Comment("Formula. Recalculate in Excel/LibreOffice.", "Finder")
    cover.column_dimensions["A"].width = 32
    cover.column_dimensions["B"].width = 28
    cover.column_dimensions["C"].width = 55
    for row in cover.iter_rows(min_row=1, max_row=26, max_col=6):
        for cell in row:
            if cell.font is None or cell.font.name != "Arial":
                cell.font = Font(name="Arial", size=11, bold=cell.font.bold if cell.font else False, color=cell.font.color if cell.font else None)

    # Courses
    ws = wb.create_sheet("Courses")
    headers = [
        "id",
        "title",
        "level",
        "subjectFamily",
        "campus",
        "durationYears",
        "tuitionGbpYearOne",
        "startDates",
        "indiaMinPercent",
        "sourceUrl",
        "checkedOn",
        "issues",
    ]
    ws.append(headers)
    ws.append(["Catalogue the app reads. Tuition is GBP for 2026/27 international full-time year one."])
    ws.merge_cells("A2:L2")
    for course in courses:
        ws.append(
            [
                course["id"],
                course["title"],
                course["level"],
                course["subjectFamily"],
                course["campus"],
                course["durationYears"],
                course["tuitionGbpYearOne"],
                ",".join(course["startDates"]),
                course["indiaEntry"]["minPercent"],
                course["sourceUrl"],
                course["checkedOn"],
                " | ".join(course["issues"]),
            ]
        )
    style_header(ws, {get_column_letter(i): w for i, w in enumerate([28, 42, 22, 16, 28, 14, 18, 24, 16, 55, 14, 60], 1)})
    for row in ws.iter_rows(min_row=3, max_row=2 + len(courses), min_col=7, max_col=7):
        for cell in row:
            cell.number_format = "£#,##0"
            cell.font = blue
    ws.row_dimensions[2].height = 18

    # Scholarships
    sh = wb.create_sheet("Scholarships")
    sh.append(["field", "value", "url", "checkedOn"])
    sh.append(["id", "uel-international-scholarship", SCHOLARSHIP, CHECKED_ON])
    sh.append(["amountMinGbp", 500, SCHOLARSHIP, CHECKED_ON])
    sh.append(["amountMaxGbp", 5000, SCHOLARSHIP, CHECKED_ON])
    sh.append(["appliesTo", "first-year-tuition", SCHOLARSHIP, CHECKED_ON])
    sh.append(["autoAssessed", "TRUE", SCHOLARSHIP, CHECKED_ON])
    sh.append(["requiresOffer", "TRUE", SCHOLARSHIP, CHECKED_ON])
    sh.append(["rangeLabel", '=CONCATENATE("£",B3,"–£",TEXT(B4,"#,##0")," off first-year tuition")', SCHOLARSHIP, CHECKED_ON])
    style_header(sh, {"A": 22, "B": 55, "C": 70, "D": 14})
    sh["B3"].number_format = "£#,##0"
    sh["B4"].number_format = "£#,##0"
    sh["B3"].font = blue
    sh["B4"].font = blue
    sh["B8"].font = Font(name="Arial", size=11, color="008000")

    # Costs
    cs = wb.create_sheet("Costs")
    cs.append(["line", "who", "valueGbp", "url", "checkedOn", "notes"])
    cost_lines = [
        ("londonLivingGbp", "Ambrosh GOV.UK", living["londonLivingGbp"], GOV_MONEY, "Spot-check London living / maintenance."),
        ("visaGbp", "Ambrosh GOV.UK", living["visaGbp"], GOV_VISA, "Spot-check student visa fee."),
        ("ihsGbp", "Ambrosh GOV.UK", living["ihsGbp"], GOV_IHS, "Spot-check Immigration Health Surcharge."),
        ("depositGbp", "Build", living["depositGbp"], DEPOSIT, "£1,500 starts pre-CAS; full amount is in the offer letter."),
        ("accommodationGbp", "Build", living["accommodationGbp"], ACCOM, "East standard full academic year 2026/27."),
        ("flightsGbp", "Build", living["flightsGbp"], "PLACEHOLDER", "Unsourced. Keep PLACEHOLDER."),
        ("inrPerGbp", "Build RBI", living["inrPerGbp"], RBI, "Reference rate as at 1.00pm 28 Oct 2026 (FBIL on RBI homepage)."),
    ]
    for line in cost_lines:
        cs.append([line[0], line[1], line[2], line[3], CHECKED_ON, line[4]])
    style_header(cs, {"A": 22, "B": 18, "C": 16, "D": 70, "E": 14, "F": 55})
    for row_idx in range(2, 9):
        cell = cs[f"C{row_idx}"]
        if cell.value == "PLACEHOLDER":
            cell.fill = yellow
            cell.font = blue
        else:
            cell.font = blue
            if isinstance(cell.value, int):
                cell.number_format = "£#,##0"
    cs["C8"].number_format = "0.0000"
    cs["A11"] = "Placeholder count"
    cs["B11"] = '=COUNTIF(C2:C7,"PLACEHOLDER")'
    cs["B11"].font = Font(name="Arial", size=11, color="008000")

    # Sources
    src = wb.create_sheet("Sources")
    src.append(["field", "value", "url", "checkedOn", "collector", "issues"])
    for row in rows:
        src.append([row["field"], row["value"], row["url"], row["checkedOn"], row["collector"], row["issues"]])
    style_header(src, {"A": 42, "B": 36, "C": 70, "D": 14, "E": 18, "F": 60})

    # GOV.UK
    gov = wb.create_sheet("GOV.UK")
    gov.append(["fact", "valueGbp", "url", "checkedOn", "who"])
    gov.append(["Student visa fee", "PLACEHOLDER", GOV_VISA, "", "Ambrosh"])
    gov.append(["Immigration Health Surcharge", "PLACEHOLDER", GOV_IHS, "", "Ambrosh"])
    gov.append(["London living / maintenance", "PLACEHOLDER", GOV_MONEY, "", "Ambrosh"])
    gov["A6"] = "Paste the figure, the page URL, and the date you checked. Build will copy them into costs.json and sources.csv."
    gov.merge_cells("A6:E7")
    style_header(gov, {"A": 36, "B": 16, "C": 70, "D": 14, "E": 14})
    for row_idx in range(2, 5):
        for col in ("B", "C", "D"):
            gov[f"{col}{row_idx}"].fill = yellow
            gov[f"{col}{row_idx}"].font = blue

    # Issues
    iss = wb.create_sheet("Issues")
    iss.append(["courseOrTopic", "issue"])
    seen = set()
    for course in COURSES:
        for issue in course["issues"]:
            key = (course["id"], issue)
            if key not in seen:
                iss.append([course["id"], issue])
                seen.add(key)
    iss.append(["accommodation", "Intro/FAQ from £189/week vs West standard table £182/week."])
    iss.append(["india-page", "Staff bio contained stray copy (“this is mine”) on the live page. Not used as a fact."])
    iss.append(["flights", "India → London fare not sourced."])
    style_header(iss, {"A": 36, "B": 100})

    out = CATALOGUE / "task2-sources.xlsx"
    wb.save(out)


def main() -> None:
    courses, _policies, living = write_json()
    rows = source_rows(courses)
    write_csv(rows)
    write_note()
    write_xlsx(courses, rows, living)
    print(f"Wrote {len(courses)} courses, {len(rows)} source rows.")


if __name__ == "__main__":
    main()
